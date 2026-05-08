"""収支管理シート操作の共通ユーティリティ。

実 Excel 受領後に拡充する。現時点ではバックアップとシート列挙の最低限のみ。
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path


def archive_backup(src: Path, archive_root: Path) -> Path:
    """編集前バックアップを `archive_root/<YYYY-MM-DD>/` 以下に作成する。"""
    if not src.is_file():
        raise FileNotFoundError(src)
    today = datetime.now().strftime("%Y-%m-%d")
    dest_dir = archive_root / today
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / src.name
    shutil.copy2(src, dest)
    return dest


def list_sheet_names(xlsx_path: Path) -> list[str]:
    """シート名を列挙。半角スペース混入を検出するため repr 込みで返す呼び出し側を想定。

    openpyxl が未インストールの環境ではインポートエラーを raise する。
    """
    from openpyxl import load_workbook  # type: ignore[import-not-found]

    wb = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
